"""Measure real XLSX exports in fresh processes with native peak-memory counters.

Only temporary synthetic state is used. The default 50,000-row product cap is
checked separately; 100k/200k measurements explicitly configure the existing
server-owned service cap, not the application/UI default. No packages installed.
"""
import argparse
import gc
import hashlib
import importlib.metadata
import json
import math
import os
from pathlib import Path
import platform
import statistics
import subprocess
import sys
import tempfile
import time
from types import SimpleNamespace


FIELDS = ('id', 'version', 'Material_Type', 'Vendor_Code', 'Vendor_Name',
          'Country', 'City', 'Rev', 'Supplier_Level')
SPECIAL = ('=1+1', '+001', '-001', '@SUM(A1)', '#N/A', '00123', '2026-10-04',
           '2026-10-04T12:34:56', 'Synthetic \u53f0\u5317', 'line1\nline2', 'A&B<text>')


def fixture_row(index):
    return dict(Material_Type='Synthetic XLSX fixture',
                Vendor_Code='SYN-{:07d}'.format(index),
                Vendor_Name='Synthetic supplier {:03d}'.format(index % 100),
                Country=('TW', 'JP', 'DE', 'US')[index % 4],
                City='Synthetic city {:02d}'.format(index % 32),
                Rev=SPECIAL[index % len(SPECIAL)], Supplier_Level='LEVEL 1')


def expected_row(index):
    row = fixture_row(index)
    row.update(id=index + 1, version=1, Supplier_Level='LEVEL 1(KEY SUPPLIER)')
    return tuple(row[key] for key in FIELDS)


def hash_row(digest, row):
    digest.update(json.dumps(row, ensure_ascii=False, separators=(',', ':')).encode('utf-8'))
    digest.update(b'\n')


def native_memory():
    if sys.platform == 'win32':
        import ctypes
        from ctypes import wintypes

        class Counters(ctypes.Structure):
            _fields_ = [('cb', wintypes.DWORD), ('PageFaultCount', wintypes.DWORD)] + [
                (name, ctypes.c_size_t) for name in ('PeakWorkingSetSize', 'WorkingSetSize',
                 'QuotaPeakPagedPoolUsage', 'QuotaPagedPoolUsage', 'QuotaPeakNonPagedPoolUsage',
                 'QuotaNonPagedPoolUsage', 'PagefileUsage', 'PeakPagefileUsage', 'PrivateUsage')]

        kernel = ctypes.WinDLL('kernel32', use_last_error=True)
        psapi = ctypes.WinDLL('psapi', use_last_error=True)
        kernel.GetCurrentProcess.restype = wintypes.HANDLE
        psapi.GetProcessMemoryInfo.argtypes = [wintypes.HANDLE, ctypes.POINTER(Counters), wintypes.DWORD]
        psapi.GetProcessMemoryInfo.restype = wintypes.BOOL
        values = Counters()
        values.cb = ctypes.sizeof(values)
        if not psapi.GetProcessMemoryInfo(kernel.GetCurrentProcess(), ctypes.byref(values), values.cb):
            raise ctypes.WinError(ctypes.get_last_error())
        return {'method': 'Windows GetProcessMemoryInfo; bytes; process-lifetime high-water marks',
                'peak_rss_bytes': values.PeakWorkingSetSize, 'rss_bytes': values.WorkingSetSize,
                'peak_commit_bytes': values.PeakPagefileUsage, 'private_bytes': values.PrivateUsage}
    import resource
    maximum = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    factor = 1 if sys.platform == 'darwin' else 1024
    return {'method': 'getrusage ru_maxrss; bytes; process-lifetime high-water mark',
            'peak_rss_bytes': int(maximum * factor), 'rss_bytes': None,
            'peak_commit_bytes': None, 'private_bytes': None}


def service_for(source, database, cap=50000):
    sys.path.insert(0, str(source))
    from reporting_workspace.legacy_crud import LegacyCrudService
    from reporting_workspace.legacy_policy import Policy
    actor = SimpleNamespace(id='synthetic-benchmark-owner', orgcode='SYNTHETIC-A',
                            is_authenticated=True, is_dev=True, is_admin=False)
    service = LegacyCrudService(str(database), lambda actor: Policy(
        orgcode=['SYNTHETIC'], crud_roles=['dev']), allow_upload=True, max_export_rows=cap)
    return service, actor


def child(mode, request):
    source, database = Path(request['source']), Path(request['database'])
    size = request['rows']
    if mode == 'seed':
        service, actor = service_for(source, database)
        digest = hashlib.sha256()
        try:
            start = time.perf_counter_ns()
            for offset in range(0, size, 500):
                batch = [fixture_row(index) for index in range(offset, min(offset + 500, size))]
                imported = service.submit_stage(actor, service.stage_rows(actor, batch))
                assert imported.committed and imported.failed == 0 and imported.created == len(batch)
                for index in range(offset, min(offset + 500, size)):
                    hash_row(digest, expected_row(index))
            return {'rows': size, 'normalized_fixture_sha256': digest.hexdigest(),
                    'setup_ms': (time.perf_counter_ns() - start) / 1000000,
                    'database_bytes': database.stat().st_size}
        finally:
            service.close()
    if mode in ('export', 'default-cap'):
        import openpyxl  # Identical module warmup before every measured process.
        service, actor = service_for(source, database, size if mode == 'export' else 50000)
        try:
            gc.collect()
            memory_before = native_memory()
            start = time.perf_counter_ns()
            if mode == 'default-cap':
                from reporting_workspace.legacy_crud import InvalidInput
                try:
                    service.export_xlsx(actor)
                except InvalidInput:
                    return {'status': 'expected-rejection', 'default_cap': 50000,
                            'wall_ms': (time.perf_counter_ns() - start) / 1000000}
                raise AssertionError('Default cap did not reject oversized export')
            content = service.export_xlsx(actor)
            elapsed = (time.perf_counter_ns() - start) / 1000000
            memory_after = native_memory()
            Path(request['xlsx']).write_bytes(content)
            return {'wall_ms': elapsed, 'memory_before': memory_before,
                    'memory_after': memory_after, 'xlsx_bytes': len(content),
                    'xlsx_sha256': hashlib.sha256(content).hexdigest(),
                    'configured_export_cap': size}
        finally:
            service.close()
    if mode == 'validate':
        import openpyxl
        workbook = openpyxl.load_workbook(request['xlsx'], read_only=True,
                                          data_only=False, keep_links=False)
        digest, count = hashlib.sha256(), 0
        try:
            assert workbook.sheetnames == ['QSL']
            iterator = workbook['QSL'].iter_rows()
            assert tuple(cell.value for cell in next(iterator)) == FIELDS
            for index, cells in enumerate(iterator):
                values = tuple(cell.value for cell in cells)
                assert values == expected_row(index), 'Exported values/order mismatch'
                assert all(cell.data_type == 'n' for cell in cells[:2])
                assert all(cell.data_type == 's' for cell in cells[2:])
                assert all(isinstance(value, int) and not isinstance(value, bool) for value in values[:2])
                assert all(isinstance(value, str) for value in values[2:])
                hash_row(digest, values)
                count += 1
            assert count == size
            assert digest.hexdigest() == request['fixture_sha256']
            return {'rows': count, 'normalized_rows_sha256': digest.hexdigest(),
                    'sheet_names': ['QSL'], 'cell_types': ['n', 'n'] + ['s'] * 7,
                    'literal_formula_error_unicode_date_like_text': 'passed'}
        finally:
            workbook.close()
    raise ValueError('Unknown child mode')


def run_child(mode, request):
    result = subprocess.run([sys.executable, '-B', str(Path(__file__).resolve()), '--child', mode],
                            input=json.dumps(request), text=True, capture_output=True, timeout=600)
    if result.returncode or result.stderr.strip():
        # No payloads or private paths are persisted in the public evidence.
        raise RuntimeError('Benchmark child {} failed (exit {}, stderr present={})'.format(
            mode, result.returncode, bool(result.stderr.strip())))
    return json.loads(result.stdout)


def source_hashes(source):
    return {path.relative_to(source).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in sorted((source / 'reporting_workspace').rglob('*.py'))}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-root', type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument('--output', type=Path)
    parser.add_argument('--rows', type=int, nargs='+', default=[100000, 200000])
    parser.add_argument('--samples', type=int, default=3)
    parser.add_argument('--warmups', type=int, default=1)
    parser.add_argument('--label', default='candidate')
    parser.add_argument('--child', choices=('seed', 'export', 'validate', 'default-cap'))
    args = parser.parse_args()
    if args.child:
        print(json.dumps(child(args.child, json.load(sys.stdin))))
        return
    if args.output is None or args.output.exists():
        parser.error('Choose a new --output file; evidence is never overwritten')
    if args.samples < 3 or args.warmups < 1 or any(n <= 50000 or n > 500000 for n in args.rows):
        parser.error('Use >=3 samples, >=1 warmup and 50,001..500,000 rows')
    source = args.source_root.resolve()
    hashes = source_hashes(source)
    result = {'schema': 1, 'label': args.label,
              'git_head': subprocess.check_output(['git', '-C', str(source), 'rev-parse', 'HEAD'], text=True).strip(),
              'python': platform.python_version(), 'platform': platform.system() + ' ' + platform.release(),
              'machine': platform.machine(), 'packages': {name: importlib.metadata.version(name)
                for name in ('openpyxl', 'Flask', 'Werkzeug', 'dash')},
              'harness_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
              'source_hashes': hashes, 'samples': args.samples, 'warmups': args.warmups,
              'timer': 'perf_counter_ns; export call only; milliseconds',
              'memory_scope': 'native process-lifetime peak in fresh child; includes imports/service setup; measured before/after export; excludes fixture seed and validator children',
              'warmup_scope': 'separate fresh processes; warms filesystem/OS caches, not the timed Python heap',
              'p95_method': 'nearest rank; three samples => maximum; descriptive only',
              'fixture_method': 'authorized 500-row stage/import batches in separate seed child',
              'default_cap_unchanged': 50000, 'datasets': []}
    with tempfile.TemporaryDirectory(prefix='large-xlsx-benchmark-') as directory:
        for size in args.rows:
            request = {'source': str(source), 'database': str(Path(directory) / ('data-{}.sqlite'.format(size))),
                       'rows': size, 'xlsx': str(Path(directory) / 'export.xlsx')}
            setup = run_child('seed', request)
            request['fixture_sha256'] = setup['normalized_fixture_sha256']
            dataset = {'rows': size, 'setup': setup,
                       'default_cap_check': run_child('default-cap', request), 'warmups': [], 'samples': []}
            for index in range(args.warmups + args.samples):
                exported = run_child('export', request)
                exported['validation'] = run_child('validate', request)
                key = 'warmups' if index < args.warmups else 'samples'
                dataset[key].append(exported)
                print(json.dumps({'label': args.label, 'rows': size, 'phase': key,
                                  'index': len(dataset[key]), 'wall_ms': round(exported['wall_ms'], 2),
                                  'peak_rss_bytes': exported['memory_after']['peak_rss_bytes']}), flush=True)
            for key, getter in (
                    ('wall_ms', lambda sample: sample['wall_ms']),
                    ('peak_rss_bytes', lambda sample: sample['memory_after']['peak_rss_bytes']),
                    ('peak_commit_bytes', lambda sample: sample['memory_after']['peak_commit_bytes'])):
                values = [getter(sample) for sample in dataset['samples']]
                if all(value is not None for value in values):
                    dataset['median_' + key] = statistics.median(values)
                    dataset['p95_' + key] = sorted(values)[math.ceil(len(values) * .95) - 1]
            result['datasets'].append(dataset)
    if source_hashes(source) != hashes:
        raise RuntimeError('Source changed during measurement; run is invalid')
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'status': 'complete', 'output': args.output.name}), flush=True)


if __name__ == '__main__':
    main()
