"""Offline QSL compatibility service. No Oracle, LDAP, SMTP or Dash side effects.

The caller supplies an authenticated *server* user, never browser claims. A fresh
policy is resolved at every public data operation. This adapter demonstrates the
transaction contract; it does not migrate company tables. SQLite is local-host
storage. Requests/callbacks must additionally enforce transport/CSRF protection.
"""
import csv
import hashlib
import io
import json
import secrets
import re
import zipfile
import xml.etree.ElementTree as ET
import sqlite3
import threading
import time
from contextlib import contextmanager
from dataclasses import dataclass
from typing import Callable


QSL_KEY = ('Material_Type', 'Vendor_Code', 'Vendor_Name', 'Country', 'City')
QSL_FIELDS = QSL_KEY + ('Rev', 'Supplier_Level')
CSV_FIELDS = ('id', 'version') + QSL_FIELDS


class LegacyCrudError(Exception):
    """Safe domain error. Never includes SQL, tokens, row data or credentials."""
    code = 'invalid_request'


class PermissionDenied(LegacyCrudError):
    code = 'permission_denied'


class InvalidInput(LegacyCrudError):
    code = 'invalid_input'


class RecordConflict(LegacyCrudError):
    code = 'record_conflict'


class RecordNotFound(LegacyCrudError):
    code = 'record_not_found'


class InvalidStage(LegacyCrudError):
    code = 'invalid_stage'


class InvalidPreview(LegacyCrudError):
    code = 'invalid_preview'


class AdapterUnavailable(LegacyCrudError):
    code = 'adapter_unavailable'


class StorageUnavailable(LegacyCrudError):
    code = 'storage_unavailable'


@dataclass(frozen=True)
class ImportResult:
    created: int = 0
    updated: int = 0
    failed: int = 0
    committed: bool = False
    dry_run: bool = False
    would_create: int = 0
    would_update: int = 0
    errors: tuple = ()
    refresh_failed: bool = False
    rows: tuple = ()


@dataclass(frozen=True)
class MutationResult:
    record_id: int
    version: int
    committed: bool
    dry_run: bool = False
    refresh_failed: bool = False
    rows: tuple = ()


def _positive_int(value):
    if isinstance(value, bool):
        raise InvalidInput('A positive integer is required.')
    if isinstance(value, str) and value.isascii() and value.isdigit():
        value = int(value) if len(value) <= 19 else 0
    if not isinstance(value, int) or not 0 < value <= 9223372036854775807:
        raise InvalidInput('A positive integer is required.')
    return value


def _text(value, maximum=255, required=False):
    if not isinstance(value, str) or len(value) > maximum or '\x00' in value:
        raise InvalidInput('Invalid field value.')
    value = value.strip()
    if required and not value:
        raise InvalidInput('Required field missing.')
    return value


def _level(value):
    value = _text(value)
    aliases = {'LEVEL 1': 'LEVEL 1(KEY SUPPLIER)',
               'LEVEL 1 (KEY SUPPLIER)': 'LEVEL 1(KEY SUPPLIER)',
               'LEVEL 1(KEY SUPPLIER)': 'LEVEL 1(KEY SUPPLIER)'}
    return aliases.get(value.upper(), value)


def _spreadsheet_safe(value):
    # Formula protection is intentionally presentation-only; do not rewrite DB.
    value = str(value)
    if value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n')):
        return "'" + value
    return value


class LegacyCrudService:
    """Server-authorized synthetic QSL service and one-use upload staging.

    policy_resolver(user) must return the current server-owned Policy on every
    call. `only_update` and `allow_upload` are server configuration, not UI state.
    Locks name fields that may not be client-supplied on that operation.
    CSV update rows require both id and version, and may not change update locks.
    `refresh` is an optional post-commit server callable; its failure is separate
    from the mutation outcome and never rolls back or retries a committed write.
    """
    def __init__(self, database, policy_resolver: Callable, target='synthetic-qsl',
                 only_update=False, allow_upload=False, create_locks=(),
                 update_locks=(), defaults=None, max_rows=50000,
                 max_file_bytes=10 * 1024 * 1024, max_export_rows=50000,
                 staging_ttl=7200, clock=time.time, preview_ttl=900):
        if not callable(policy_resolver):
            raise TypeError('A server policy resolver is required.')
        self.target = _text(target, 128, True)
        self.policy_resolver = policy_resolver
        self.only_update = bool(only_update)
        self.allow_upload = bool(allow_upload)
        self.create_locks = frozenset(create_locks)
        self.update_locks = frozenset(update_locks)
        if not (self.create_locks | self.update_locks).issubset(QSL_FIELDS):
            raise InvalidInput('Unknown locked field.')
        self.defaults = dict(defaults or {})
        if not set(self.defaults).issubset(QSL_FIELDS):
            raise InvalidInput('Unknown default field.')
        self.max_rows = _positive_int(max_rows)
        self.max_file_bytes = _positive_int(max_file_bytes)
        self.max_export_rows = _positive_int(max_export_rows)
        self.staging_ttl = _positive_int(staging_ttl)
        self.preview_ttl = _positive_int(preview_ttl)
        self.clock = clock
        self._lock = threading.RLock()
        self._db = sqlite3.connect(database, timeout=10, check_same_thread=False,
                                   isolation_level=None)
        self._db.row_factory = sqlite3.Row
        try:
            self._initialize()
        except Exception:
            # Constructor failure must not retain a Windows file handle until GC.
            self._db.close()
            raise

    def _initialize(self):
        self._db.executescript('''
          CREATE TABLE IF NOT EXISTS legacy_qsl_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT, target TEXT NOT NULL,
            Material_Type TEXT NOT NULL, Vendor_Code TEXT NOT NULL,
            Vendor_Name TEXT NOT NULL, Country TEXT NOT NULL, City TEXT NOT NULL,
            Rev TEXT NOT NULL, Supplier_Level TEXT NOT NULL,
            version INTEGER NOT NULL DEFAULT 1, deleted INTEGER NOT NULL DEFAULT 0);
          CREATE UNIQUE INDEX IF NOT EXISTS legacy_qsl_unique_active ON
            legacy_qsl_records(target, Material_Type, Vendor_Code, Vendor_Name,
              Country, City) WHERE deleted=0;
          CREATE TABLE IF NOT EXISTS legacy_qsl_audit (
            id INTEGER PRIMARY KEY AUTOINCREMENT, target TEXT NOT NULL,
            actor TEXT NOT NULL, action TEXT NOT NULL, record_id INTEGER NOT NULL,
            version INTEGER NOT NULL, changed_fields TEXT NOT NULL, at REAL NOT NULL);
          CREATE TABLE IF NOT EXISTS legacy_qsl_stages (
            token_hash TEXT PRIMARY KEY, owner TEXT NOT NULL, target TEXT NOT NULL,
            expires REAL NOT NULL, consumed INTEGER NOT NULL DEFAULT 0,
            payload TEXT NOT NULL);
          CREATE TABLE IF NOT EXISTS legacy_qsl_revisions (
            target TEXT NOT NULL, record_id INTEGER NOT NULL,
            version INTEGER NOT NULL, actor TEXT NOT NULL, action TEXT NOT NULL,
            at REAL NOT NULL, changed_fields TEXT NOT NULL,
            source_version INTEGER, snapshot TEXT NOT NULL,
            PRIMARY KEY(target,record_id,version));
          CREATE TRIGGER IF NOT EXISTS legacy_qsl_revisions_no_update
            BEFORE UPDATE ON legacy_qsl_revisions BEGIN
              SELECT RAISE(ABORT, 'Revision history is immutable.');
            END;
          CREATE TRIGGER IF NOT EXISTS legacy_qsl_revisions_no_delete
            BEFORE DELETE ON legacy_qsl_revisions BEGIN
              SELECT RAISE(ABORT, 'Revision history is immutable.');
            END;
          CREATE TABLE IF NOT EXISTS legacy_qsl_previews (
            token_hash TEXT PRIMARY KEY, owner TEXT NOT NULL, target TEXT NOT NULL,
            action TEXT NOT NULL, record_id INTEGER NOT NULL,
            expected_version INTEGER NOT NULL, source_version INTEGER,
            expires REAL NOT NULL, consumed INTEGER NOT NULL DEFAULT 0,
            payload TEXT NOT NULL);
        ''')
        # Old databases contain only current rows. Preserve that known snapshot,
        # never manufacture prior versions whose values cannot be reconstructed.
        with self._transaction():
            rows = self._db.execute('''SELECT r.* FROM legacy_qsl_records r
                WHERE r.target=? AND NOT EXISTS (
                  SELECT 1 FROM legacy_qsl_revisions h WHERE h.target=r.target
                  AND h.record_id=r.id AND h.version=r.version)''',
                (self.target,)).fetchall()
            for row in rows:
                audit = self._db.execute('''SELECT actor,at FROM legacy_qsl_audit
                    WHERE target=? AND record_id=? AND version=?
                    ORDER BY id DESC LIMIT 1''',
                    (self.target, row['id'], row['version'])).fetchone()
                self._save_revision(row, audit['actor'] if audit else 'legacy-baseline',
                                    'baseline', audit['at'] if audit else self.clock(), ())

    def close(self):
        with self._lock:
            self._db.close()

    def _authorize(self, user, action):
        from .legacy_policy import authorize, identity_id
        if action not in ('read', 'export', 'create', 'update', 'delete', 'restore', 'upload', 'audit'):
            raise PermissionDenied('Operation denied.')
        if self.only_update and action in ('create', 'delete', 'restore'):
            raise PermissionDenied('Operation denied.')
        if action == 'upload' and not self.allow_upload:
            raise PermissionDenied('Operation denied.')
        access = 'read' if action in ('read', 'export') else 'crud'
        try:
            permitted = authorize(user, self.policy_resolver(user), access)
            actor = identity_id(getattr(user, 'id', None))
        except Exception:
            raise PermissionDenied('Operation denied.') from None
        if not permitted:
            raise PermissionDenied('Operation denied.')
        return actor

    @contextmanager
    def _transaction(self, dry_run=False):
        if not isinstance(dry_run, bool):
            raise InvalidInput('Invalid dry-run option.')
        with self._lock:
            try:
                self._db.execute('BEGIN IMMEDIATE')
                yield
                self._db.execute('ROLLBACK' if dry_run else 'COMMIT')
            except Exception as exc:
                if self._db.in_transaction:
                    self._db.execute('ROLLBACK')
                if isinstance(exc, LegacyCrudError):
                    raise
                if isinstance(exc, sqlite3.IntegrityError):
                    raise RecordConflict('Record conflicts with current data.') from None
                if isinstance(exc, sqlite3.Error):
                    raise StorageUnavailable('Storage operation failed.') from None
                raise

    def _payload(self, source, updating=False, existing=None):
        if not isinstance(source, dict) or not set(source).issubset(QSL_FIELDS):
            raise InvalidInput('Unknown or invalid fields.')
        locks = self.update_locks if updating else self.create_locks
        payload = {}
        for name, value in source.items():
            value = _level(value) if name == 'Supplier_Level' else _text(value, required=name in QSL_KEY)
            if name in locks:
                # Full CSV snapshots may carry a locked field only unchanged.
                if not updating or existing is None or value != existing[name]:
                    raise InvalidInput('A locked field cannot be changed.')
                continue
            payload[name] = value
        if not updating:
            base = dict(self.defaults)
            base.update(payload)
            for name in QSL_FIELDS:
                value = base.get(name, '')
                payload[name] = _level(value) if name == 'Supplier_Level' else _text(value, required=name in QSL_KEY)
        return payload

    def _record(self, record_id, deleted=0):
        record_id = _positive_int(record_id)
        row = self._db.execute('SELECT * FROM legacy_qsl_records WHERE id=? AND target=? AND deleted=?',
                               (record_id, self.target, deleted)).fetchone()
        if row is None:
            raise RecordNotFound('Record unavailable.')
        return row

    def _record_any(self, record_id):
        row = self._db.execute('SELECT * FROM legacy_qsl_records WHERE id=? AND target=?',
                               (_positive_int(record_id), self.target)).fetchone()
        if row is None:
            raise RecordNotFound('Record unavailable.')
        return row

    @staticmethod
    def _snapshot(row):
        return {name: row[name] for name in CSV_FIELDS + ('deleted',)}

    def _save_revision(self, row, actor, action, at, fields, source_version=None):
        self._db.execute('''INSERT INTO legacy_qsl_revisions
            (target,record_id,version,actor,action,at,changed_fields,source_version,snapshot)
            VALUES(?,?,?,?,?,?,?,?,?)''',
            (self.target, row['id'], row['version'], actor, action, at,
             json.dumps(sorted(fields)), source_version,
             json.dumps(self._snapshot(row), ensure_ascii=False, separators=(',', ':'))))

    def _audit(self, actor, action, record_id, version, fields, source_version=None):
        at = self.clock()
        self._db.execute('INSERT INTO legacy_qsl_audit(target,actor,action,record_id,version,changed_fields,at) VALUES(?,?,?,?,?,?,?)',
                         (self.target, actor, action, record_id, version,
                          json.dumps(sorted(fields)), at))
        self._save_revision(self._record_any(record_id), actor, action, at,
                            fields, source_version)

    def _create(self, actor, payload):
        if self.only_update:
            raise PermissionDenied('Creation disabled.')
        data = self._payload(payload)
        fields = ','.join(QSL_FIELDS)
        result = self._db.execute('INSERT INTO legacy_qsl_records(target,' + fields + ') VALUES(' + ','.join(['?'] * 8) + ')',
                                  [self.target] + [data[name] for name in QSL_FIELDS])
        self._audit(actor, 'create', result.lastrowid, 1, data)
        return result.lastrowid, 1

    def _update(self, actor, record_id, version, payload):
        version = _positive_int(version)
        row = self._record(record_id)
        if row['version'] != version or version == 9223372036854775807:
            raise RecordConflict('Record version changed.')
        data = self._payload(payload, updating=True, existing=row)
        if not data:
            raise InvalidInput('No editable fields supplied.')
        statement = ','.join(name + '=?' for name in data)
        changed = [name for name in data if row[name] != data[name]]
        self._db.execute('UPDATE legacy_qsl_records SET ' + statement + ', version=version+1 WHERE id=? AND target=? AND version=?',
                         list(data.values()) + [row['id'], self.target, version])
        self._audit(actor, 'update', row['id'], version + 1, changed)
        return row['id'], version + 1

    def _finish(self, rid, version, dry_run, refresh):
        rows, failed = (), False
        if refresh is not None and not dry_run:
            try:
                rows = tuple(refresh())
            except Exception:
                failed = True
        return MutationResult(rid, version, not dry_run, dry_run, failed, rows)

    def create(self, user, payload, dry_run=False, refresh=None):
        actor = self._authorize(user, 'create')
        with self._transaction(dry_run):
            rid, version = self._create(actor, payload)
        return self._finish(rid, version, dry_run, refresh)

    def update(self, user, record_id, version, payload, dry_run=False, refresh=None):
        actor = self._authorize(user, 'update')
        with self._transaction(dry_run):
            rid, version = self._update(actor, record_id, version, payload)
        return self._finish(rid, version, dry_run, refresh)

    def _lifecycle(self, user, record_id, version, restore, dry_run, refresh):
        actor = self._authorize(user, 'restore' if restore else 'delete')
        version = _positive_int(version)
        with self._transaction(dry_run):
            row = self._record(record_id, deleted=1 if restore else 0)
            if row['version'] != version or version == 9223372036854775807:
                raise RecordConflict('Record version changed.')
            self._db.execute('UPDATE legacy_qsl_records SET deleted=?,version=version+1 WHERE id=? AND target=?',
                             (0 if restore else 1, row['id'], self.target))
            self._audit(actor, 'restore' if restore else 'delete', row['id'], version + 1, ('deleted',))
        return self._finish(row['id'], version + 1, dry_run, refresh)

    def delete(self, user, record_id, version, dry_run=False, refresh=None):
        """Recoverable soft deletion; there is deliberately no purge API."""
        return self._lifecycle(user, record_id, version, False, dry_run, refresh)

    def restore(self, user, record_id, version, dry_run=False, refresh=None):
        return self._lifecycle(user, record_id, version, True, dry_run, refresh)

    @staticmethod
    def _check_version(row, expected_version):
        expected_version = _positive_int(expected_version)
        if row['version'] != expected_version or expected_version == 9223372036854775807:
            raise RecordConflict('Record version changed.')
        return expected_version

    @staticmethod
    def _diff(before, after):
        # Automatic version increments are metadata, not user-editable changes.
        return [{'field': name, 'before': before[name], 'after': after[name]}
                for name in QSL_FIELDS + ('deleted',) if before[name] != after[name]]

    def _check_unique(self, snapshot):
        if snapshot['deleted']:
            return
        clauses = ' AND '.join(name + '=?' for name in QSL_KEY)
        duplicate = self._db.execute('''SELECT id FROM legacy_qsl_records
            WHERE target=? AND deleted=0 AND id<>? AND ''' + clauses,
            [self.target, snapshot['id']] + [snapshot[name] for name in QSL_KEY]).fetchone()
        if duplicate is not None:
            raise RecordConflict('Record conflicts with current data.')

    def _save_preview(self, actor, action, before, after, source_version=None):
        diff = self._diff(before, after)
        if not diff:
            raise InvalidInput('No changes to confirm.')
        self._check_unique(after)
        token = secrets.token_urlsafe(32)
        expires = self.clock() + self.preview_ttl
        payload = {name: after[name] for name in QSL_FIELDS + ('deleted',)}
        self._db.execute('''INSERT INTO legacy_qsl_previews
            (token_hash,owner,target,action,record_id,expected_version,source_version,expires,payload)
            VALUES(?,?,?,?,?,?,?,?,?)''',
            (self._token_hash(token), actor, self.target, action, before['id'],
             before['version'], source_version, expires,
             json.dumps(payload, ensure_ascii=False, separators=(',', ':'))))
        return {'token': token, 'action': action, 'record_id': before['id'],
                'expected_version': before['version'], 'source_version': source_version,
                'expires_at': expires, 'before': before, 'after': after, 'diff': diff}

    def preview_update(self, user, record_id, payload, expected_version):
        """Validate and stage a server-owned diff; this does not modify a record.

        The opaque one-use token binds actor, target, record, normalized values,
        expected version, operation and expiry. Browser diff/Store values are
        display-only and are never accepted by confirmation.
        """
        self._authorize(user, 'update')
        with self._transaction():
            actor = self._authorize(user, 'update')
            row = self._record(record_id)
            version = self._check_version(row, expected_version)
            data = self._payload(payload, updating=True, existing=row)
            before = self._snapshot(row)
            after = dict(before, version=version + 1)
            after.update(data)
            return self._save_preview(actor, 'update', before, after)

    def _historical_snapshot(self, record_id, version):
        historical = self._db.execute('''SELECT snapshot FROM legacy_qsl_revisions
            WHERE target=? AND record_id=? AND version=?''',
            (self.target, _positive_int(record_id), _positive_int(version))).fetchone()
        if historical is None:
            raise RecordNotFound('Revision unavailable.')
        return json.loads(historical['snapshot'])

    def preview_restore(self, user, record_id, history_version, expected_version):
        """Stage a historical snapshot as a new version, including deleted state.

        No history is rewritten. A pre-upgrade record has only its known current
        baseline; unavailable earlier versions cannot be selected.
        """
        self._authorize(user, 'restore')
        with self._transaction():
            actor = self._authorize(user, 'restore')
            row = self._record_any(record_id)
            version = self._check_version(row, expected_version)
            history_version = _positive_int(history_version)
            historical = self._historical_snapshot(row['id'], history_version)
            data = self._payload({name: historical[name] for name in QSL_FIELDS},
                                 updating=True, existing=row)
            before = self._snapshot(row)
            after = dict(before, version=version + 1, deleted=historical['deleted'])
            after.update(data)
            return self._save_preview(actor, 'restore_version', before, after, history_version)

    def _confirm_preview(self, user, token, restore_only, refresh):
        actor = self._authorize(user, 'restore' if restore_only else 'update')
        try:
            digest = self._token_hash(token)
        except InvalidStage:
            raise InvalidPreview('Change preview unavailable.') from None
        with self._transaction():
            preview = self._db.execute('''SELECT * FROM legacy_qsl_previews
                WHERE token_hash=? AND owner=? AND target=? AND consumed=0 AND expires>?''',
                (digest, actor, self.target, self.clock())).fetchone()
            if preview is None or preview['action'] not in ('update', 'restore_version'):
                raise InvalidPreview('Change preview unavailable.')
            restoring = preview['action'] == 'restore_version'
            if restore_only and not restoring:
                raise InvalidPreview('Change preview unavailable.')
            # Re-resolve the operation-specific policy after acquiring the write
            # transaction: preview-time or browser permissions never authorize.
            actor = self._authorize(user, 'restore' if restoring else 'update')
            if actor != preview['owner'] or preview['expires'] <= self.clock():
                raise InvalidPreview('Change preview unavailable.')
            row = self._record_any(preview['record_id'])
            version = self._check_version(row, preview['expected_version'])
            payload = json.loads(preview['payload'])
            data = self._payload({name: payload[name] for name in QSL_FIELDS},
                                 updating=True, existing=row)
            if restoring:
                historical = self._historical_snapshot(row['id'], preview['source_version'])
                if any(historical[name] != payload[name] for name in QSL_FIELDS + ('deleted',)):
                    raise InvalidPreview('Change preview unavailable.')
                statement = ','.join([name + '=?' for name in data] + ['deleted=?', 'version=version+1'])
                self._db.execute('UPDATE legacy_qsl_records SET ' + statement +
                                 ' WHERE id=? AND target=? AND version=?',
                                 list(data.values()) + [payload['deleted'], row['id'], self.target, version])
                fields = [change['field'] for change in
                          self._diff(self._snapshot(row), dict(self._snapshot(row), **payload))]
                self._audit(actor, 'restore_version', row['id'], version + 1,
                            fields, preview['source_version'])
                rid, version = row['id'], version + 1
            else:
                if row['deleted'] or payload['deleted'] != row['deleted']:
                    raise RecordConflict('Record version changed.')
                rid, version = self._update(actor, row['id'], version, data)
            # Mutation, immutable snapshot, audit and token consumption commit
            # together. Failure rolls back all four; stale previews cannot win.
            self._db.execute("UPDATE legacy_qsl_previews SET consumed=1,payload='{}' WHERE token_hash=?",
                             (digest,))
        return self._finish(rid, version, False, refresh)

    def confirm_update(self, user, token, refresh=None):
        """Confirm either kind of staged change using only its opaque token."""
        return self._confirm_preview(user, token, False, refresh)

    def confirm_restore(self, user, token, refresh=None):
        """Restore-only confirmation; update previews cannot be used here."""
        return self._confirm_preview(user, token, True, refresh)

    def discard_preview(self, user, token):
        """Cancel a current user's staged change without changing the record."""
        actor = self._authorize(user, 'audit')
        try:
            digest = self._token_hash(token)
        except InvalidStage:
            raise InvalidPreview('Change preview unavailable.') from None
        with self._transaction():
            result = self._db.execute("""UPDATE legacy_qsl_previews SET consumed=1,payload='{}'
                WHERE token_hash=? AND owner=? AND target=? AND consumed=0""",
                (digest, actor, self.target))
            if result.rowcount != 1:
                raise InvalidPreview('Change preview unavailable.')

    def cleanup_expired_previews(self, user):
        self._authorize(user, 'audit')
        with self._transaction():
            result = self._db.execute("""UPDATE legacy_qsl_previews SET consumed=1,payload='{}'
                WHERE target=? AND consumed=0 AND expires<=?""", (self.target, self.clock()))
        return result.rowcount

    def history(self, user, record_id):
        """CRUD-authorized, target-scoped immutable snapshots, newest first."""
        self._authorize(user, 'audit')
        with self._lock:
            row = self._record_any(record_id)
            revisions = self._db.execute('''SELECT record_id,version,actor,action,at,
                changed_fields,source_version,snapshot FROM legacy_qsl_revisions
                WHERE target=? AND record_id=? ORDER BY version DESC''',
                (self.target, row['id'])).fetchall()
        result = []
        for revision in revisions:
            item = dict(revision)
            item['changed_fields'] = json.loads(item['changed_fields'])
            item['snapshot'] = json.loads(item['snapshot'])
            result.append(item)
        return result

    def get(self, user, record_id):
        """Load a current version for editing, with read authorization."""
        self._authorize(user, 'read')
        with self._lock:
            row = self._record(record_id)
        return {name: row[name] for name in CSV_FIELDS}

    def query(self, user, filters=None, limit=1000, offset=0):
        self._authorize(user, 'read')
        limit = _positive_int(limit)
        if limit > self.max_export_rows or isinstance(offset, bool) or not isinstance(offset, int) or not 0 <= offset <= 1000000:
            raise InvalidInput('Invalid query bounds.')
        filters = {} if filters is None else filters
        if not isinstance(filters, dict) or not set(filters).issubset(QSL_FIELDS):
            raise InvalidInput('Invalid filter fields.')
        clauses, values = ['target=?', 'deleted=0'], [self.target]
        for name, value in filters.items():
            # Literal substring matching, not user-controlled regex/LIKE patterns.
            clauses.append('instr(' + name + ',?)>0')
            values.append(_text(value, 200))
        with self._lock:
            rows = self._db.execute('SELECT id,version,' + ','.join(QSL_FIELDS) + ' FROM legacy_qsl_records WHERE ' +
                                    ' AND '.join(clauses) + ' ORDER BY id LIMIT ? OFFSET ?', values + [limit, offset]).fetchall()
        return [dict(row) for row in rows]

    def export_csv(self, user, filters=None):
        self._authorize(user, 'export')
        # Detect truncation rather than silently exporting an incomplete result.
        with self._transaction():
            rows = self.query(user, filters, self.max_export_rows)
            if self.query(user, filters, 1, self.max_export_rows):
                raise InvalidInput('Export limit exceeded; narrow the filters.')
        out = io.StringIO(newline='')
        writer = csv.DictWriter(out, fieldnames=CSV_FIELDS)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _spreadsheet_safe(row[key]) for key in CSV_FIELDS})
        return out.getvalue().encode('utf-8-sig')

    def template_csv(self, user):
        self._authorize(user, 'upload')
        out = io.StringIO(newline='')
        csv.writer(out).writerow(CSV_FIELDS)
        return out.getvalue().encode('utf-8-sig')

    def stage_csv(self, user, content):
        self._authorize(user, 'upload')
        if not isinstance(content, bytes) or len(content) > self.max_file_bytes:
            raise InvalidInput('Upload size limit exceeded.')
        try:
            reader = csv.DictReader(io.StringIO(content.decode('utf-8-sig'), newline=''), strict=True)
            names = reader.fieldnames or []
            if len(names) != len(set(names)) or not set(names).issubset(CSV_FIELDS) or not set(QSL_KEY).issubset(names):
                raise InvalidInput('Invalid CSV columns.')
            rows = []
            for row in reader:
                if len(rows) >= self.max_rows or None in row or any(value is None for value in row.values()):
                    raise InvalidInput('Invalid CSV shape or row limit exceeded.')
                rows.append(row)
        except (UnicodeError, csv.Error):
            raise InvalidInput('Invalid UTF-8 CSV.') from None
        return self.stage_rows(user, rows)

    def stage_xlsx(self, user, content):
        """Optional XLSX adapter, one sheet, bounded ZIP/XML and no formulas.

        The dependency must already be approved and installed by the integrator;
        this service never installs or upgrades it. Values are literal strings,
        except positive integer id/version. Unsupported typed cells are rejected.
        """
        self._authorize(user, 'upload')
        if not isinstance(content, bytes) or len(content) > self.max_file_bytes:
            raise InvalidInput('Upload size limit exceeded.')
        try:
            import openpyxl
        except ImportError:
            raise AdapterUnavailable('XLSX adapter is not installed.') from None
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or len({item.filename for item in entries}) != len(entries):
                    raise InvalidInput('Invalid XLSX package.')
                if sum(item.file_size for item in entries) > min(50 * 1024 * 1024, self.max_file_bytes * 20):
                    raise InvalidInput('Expanded XLSX limit exceeded.')
                for item in entries:
                    if item.flag_bits & 1 or item.file_size > max(1, item.compress_size) * 200:
                        raise InvalidInput('Unsafe XLSX package.')
                    if item.filename.lower().endswith('.xml'):
                        xml = archive.read(item)
                        if b'\x00' in xml or b'<!DOCTYPE' in xml.upper() or b'<!ENTITY' in xml.upper():
                            raise InvalidInput('Unsafe XLSX XML.')
                        if item.filename.startswith('xl/worksheets/'):
                            # Do not trust dimensions or ignore far-away cells.
                            for _, elem in ET.iterparse(io.BytesIO(xml), events=('end',)):
                                if elem.tag.rsplit('}', 1)[-1] == 'c':
                                    coordinate = elem.attrib.get('r', '')
                                    match = re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]{0,6})', coordinate)
                                    if not match:
                                        raise InvalidInput('Invalid XLSX coordinate.')
                                    column = 0
                                    for letter in match.group(1):
                                        column = column * 26 + ord(letter) - 64
                                    if column > len(CSV_FIELDS) or int(match.group(2)) > self.max_rows + 1:
                                        raise InvalidInput('XLSX row or column limit exceeded.')
                                elem.clear()
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True,
                                               data_only=False, keep_links=False)
            try:
                if len(workbook.worksheets) != 1:
                    raise InvalidInput('XLSX must contain exactly one worksheet.')
                sheet = workbook.worksheets[0]
                sheet.reset_dimensions()
                iterator = sheet.iter_rows(max_row=self.max_rows + 1, max_col=len(CSV_FIELDS))
                header = next(iterator, ())
                names = [cell.value for cell in header]
                while names and names[-1] is None:
                    names.pop()
                if not names or not all(isinstance(name, str) for name in names) or len(names) != len(set(names)) or not set(names).issubset(CSV_FIELDS) or not set(QSL_KEY).issubset(names):
                    raise InvalidInput('Invalid XLSX columns.')
                rows = []
                for cells in iterator:
                    if all(cell.value is None for cell in cells):
                        continue
                    if any(cell.data_type in ('f', 'e') for cell in cells):
                        raise InvalidInput('Formula and error cells are forbidden.')
                    if any(cell.value is not None for cell in cells[len(names):]):
                        raise InvalidInput('Unlabelled XLSX columns.')
                    rows.append({name: cell.value if cell.value is not None else ''
                                 for name, cell in zip(names, cells)})
            finally:
                workbook.close()
        except LegacyCrudError:
            raise
        except Exception:
            raise InvalidInput('Invalid XLSX file.') from None
        return self.stage_rows(user, rows)

    def export_xlsx(self, user, filters=None):
        self._authorize(user, 'export')
        try:
            import openpyxl
            from openpyxl.cell import WriteOnlyCell
        except ImportError:
            raise AdapterUnavailable('XLSX adapter is not installed.') from None
        with self._transaction():
            rows = self.query(user, filters, self.max_export_rows)
            if self.query(user, filters, 1, self.max_export_rows):
                raise InvalidInput('Export limit exceeded; narrow the filters.')
        workbook = openpyxl.Workbook(write_only=True)
        sheet = workbook.create_sheet('QSL')
        sheet.append(CSV_FIELDS)
        for row in rows:
            cells = []
            for name in CSV_FIELDS:
                cell = WriteOnlyCell(sheet, value=row[name])
                if name in QSL_FIELDS:
                    cell.data_type = 's'  # Literal text, even when starting '='.
                cells.append(cell)
            sheet.append(cells)
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    def stage_rows(self, user, rows):
        actor = self._authorize(user, 'upload')
        if not isinstance(rows, list) or not 0 < len(rows) <= self.max_rows:
            raise InvalidInput('Invalid upload row count.')
        for row in rows:
            if not isinstance(row, dict) or not set(row).issubset(CSV_FIELDS):
                raise InvalidInput('Invalid upload fields.')
            for key, value in row.items():
                if key in ('id', 'version'):
                    if value not in ('', None):
                        _positive_int(value)
                else:
                    _text(value, required=key in QSL_KEY)
        payload = json.dumps(rows, ensure_ascii=False, separators=(',', ':'))
        if len(payload.encode('utf-8')) > self.max_file_bytes:
            raise InvalidInput('Upload size limit exceeded.')
        token = secrets.token_urlsafe(32)
        with self._transaction():
            self._db.execute('INSERT INTO legacy_qsl_stages(token_hash,owner,target,expires,payload) VALUES(?,?,?,?,?)',
                             (self._token_hash(token), actor, self.target, self.clock() + self.staging_ttl, payload))
        return token

    @staticmethod
    def _token_hash(token):
        if not isinstance(token, str) or not 32 <= len(token) <= 128 or not all(c.isascii() and (c.isalnum() or c in '_-') for c in token):
            raise InvalidStage('Upload stage unavailable.')
        return hashlib.sha256(token.encode('ascii')).hexdigest()

    def stage_preview(self, user, token):
        """Validate the batch transactionally; return at most ten row diffs.

        A rollback-only transaction exercises the same mutation path as submit,
        including normalization, locks and conflicts within this batch. Summary
        counts cover the whole batch; no record, audit or revision is retained.
        The preview is advisory: submit rechecks authorization and current data.
        """
        actor = self._authorize(user, 'upload')
        operations, created, updated, failed = [], 0, 0, 0
        with self._transaction(dry_run=True):
            stage = self._db.execute('SELECT payload FROM legacy_qsl_stages WHERE token_hash=? AND owner=? AND target=? AND consumed=0 AND expires>?',
                                     (self._token_hash(token), actor, self.target, self.clock())).fetchone()
            if stage is None:
                raise InvalidStage('Upload stage unavailable.')
            rows = json.loads(stage['payload'])
            if len(rows) > self.max_rows or len(stage['payload'].encode('utf-8')) > self.max_file_bytes:
                raise InvalidInput('Upload limit exceeded.')
            for index, row in enumerate(rows, 1):
                operation = 'create' if row.get('id') in ('', None) else 'update'
                entry = {'row': index, 'operation': operation, 'status': 'ready',
                         'before': None, 'after': None, 'diff': [], 'error': None}
                self._db.execute('SAVEPOINT preview_row')
                try:
                    if operation == 'update':
                        entry['before'] = self._snapshot(self._record(row['id']))
                    action, rid, _ = self._apply_import_row(actor, row)
                    if action == 'create':
                        created += 1
                    else:
                        updated += 1
                    if index <= 10:
                        entry['after'] = self._snapshot(self._record_any(rid))
                        entry['diff'] = self._diff(entry['before'], entry['after']) if entry['before'] else [
                            {'field': name, 'before': None, 'after': entry['after'][name]}
                            for name in QSL_FIELDS]
                        if action == 'create':
                            # A rollback-only insert does not reserve an ID.
                            entry['after']['id'] = None
                    self._db.execute('RELEASE preview_row')
                except (LegacyCrudError, sqlite3.IntegrityError) as exc:
                    self._db.execute('ROLLBACK TO preview_row')
                    self._db.execute('RELEASE preview_row')
                    failed += 1
                    code = exc.code if isinstance(exc, LegacyCrudError) else 'record_conflict'
                    entry['status'] = 'conflict' if code == 'record_conflict' else 'invalid'
                    entry['error'] = code
                if index <= 10:
                    operations.append(entry)
        return {'rows': rows[:10], 'count': len(rows), 'operations': operations,
                'would_create': created, 'would_update': updated, 'failed': failed,
                'truncated': len(rows) > 10}

    def _apply_import_row(self, actor, row):
        rid, version = row.get('id'), row.get('version')
        data = {key: value for key, value in row.items() if key in QSL_FIELDS}
        if rid not in ('', None):
            rid, version = self._update(actor, rid, version, data)
            return 'update', rid, version
        if version not in ('', None):
            raise InvalidInput('New records cannot set a version.')
        rid, version = self._create(actor, data)
        return 'create', rid, version

    def submit_stage(self, user, token, atomic=True, dry_run=False, refresh=None):
        actor = self._authorize(user, 'upload')
        if not isinstance(atomic, bool) or not isinstance(dry_run, bool):
            raise InvalidInput('Invalid import options.')
        digest = self._token_hash(token)
        created, updated, errors = 0, 0, []
        with self._transaction(dry_run):
            stage = self._db.execute('SELECT * FROM legacy_qsl_stages WHERE token_hash=? AND owner=? AND target=? AND consumed=0 AND expires>?',
                                     (digest, actor, self.target, self.clock())).fetchone()
            if stage is None:
                raise InvalidStage('Upload stage unavailable.')
            rows = json.loads(stage['payload'])
            if len(rows) > self.max_rows or len(stage['payload'].encode('utf-8')) > self.max_file_bytes:
                raise InvalidInput('Upload limit exceeded.')
            self._db.execute('SAVEPOINT entire_import')
            for index, row in enumerate(rows, 1):
                self._db.execute('SAVEPOINT import_row')
                try:
                    action, _, _ = self._apply_import_row(actor, row)
                    if action == 'update':
                        updated += 1
                    else:
                        created += 1
                    self._db.execute('RELEASE import_row')
                except (LegacyCrudError, sqlite3.IntegrityError) as exc:
                    self._db.execute('ROLLBACK TO import_row')
                    self._db.execute('RELEASE import_row')
                    errors.append((index, exc.code if isinstance(exc, LegacyCrudError) else 'record_conflict'))
            if errors and atomic:
                self._db.execute('ROLLBACK TO entire_import')
                created, updated = 0, 0
            self._db.execute('RELEASE entire_import')
            # A rejected/partial/complete non-dry-run attempt consumes its token.
            self._db.execute("UPDATE legacy_qsl_stages SET consumed=1,payload='[]' WHERE token_hash=?", (digest,))
        success = not dry_run and (created + updated > 0)
        refreshed, refresh_failed = (), False
        if success and refresh is not None:
            try:
                refreshed = tuple(refresh())
            except Exception:
                refresh_failed = True
        return ImportResult(created=created if not dry_run else 0,
                            updated=updated if not dry_run else 0, failed=len(errors),
                            committed=success, dry_run=dry_run,
                            would_create=created if dry_run else 0,
                            would_update=updated if dry_run else 0,
                            errors=tuple(errors[:100]), refresh_failed=refresh_failed,
                            rows=refreshed)

    def discard_stage(self, user, token):
        actor = self._authorize(user, 'upload')
        with self._transaction():
            result = self._db.execute("UPDATE legacy_qsl_stages SET consumed=1,payload='[]' WHERE token_hash=? AND owner=? AND target=? AND consumed=0",
                                      (self._token_hash(token), actor, self.target))
            if result.rowcount != 1:
                raise InvalidStage('Upload stage unavailable.')

    def cleanup_expired_stages(self, user):
        """Authorized target-local maintenance; clear payloads, retain tombstones."""
        self._authorize(user, 'audit')
        with self._transaction():
            result = self._db.execute("UPDATE legacy_qsl_stages SET consumed=1,payload='[]' WHERE target=? AND consumed=0 AND expires<=?",
                                      (self.target, self.clock()))
        return result.rowcount

    def audit(self, user, limit=100):
        self._authorize(user, 'audit')
        limit = _positive_int(limit)
        if limit > 1000:
            raise InvalidInput('Audit limit exceeded.')
        with self._lock:
            rows = self._db.execute('SELECT actor,action,record_id,version,changed_fields,at FROM legacy_qsl_audit WHERE target=? ORDER BY id DESC LIMIT ?',
                                    (self.target, limit)).fetchall()
        return [dict(row) for row in rows]
