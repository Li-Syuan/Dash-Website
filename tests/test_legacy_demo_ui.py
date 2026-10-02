"""HTTP callback contract tests against the exact legacy Dash version."""
import base64
import csv
import io
import json
import re
import unittest

from reporting_workspace.legacy_demo_ui import create_demo, FIELDS, PREFIX


class DemoTransportTests(unittest.TestCase):
    def setUp(self):
        self.app = create_demo()
        self.server = self.app.server
        self.server.config['TESTING'] = True
        self.client = self.server.test_client()
        self.client.get(PREFIX)

    def tearDown(self):
        self.server.extensions['qa_demo_builder'].close()
        self.server.extensions['qa_demo_crud'].close()
        self.server.extensions['qa_demo_temporary'].cleanup()

    def login(self, who):
        page = self.client.get(PREFIX + 'demo-login').get_data(as_text=True)
        csrf = re.search('name="csrf" value="([^"]+)"', page).group(1)
        response = self.client.post(PREFIX + 'demo-login', data={'csrf': csrf, 'identity': who})
        self.assertEqual(response.status_code, 302)

    def call(self, output_prefix, trigger, values=None, headers=None):
        values = values or {}
        key = next(k for k in self.app.callback_map if output_prefix in k)
        spec = self.app.callback_map[key]
        outputs = spec['output']
        payload = {
            'output': key,
            'outputs': [{'id': item.component_id, 'property': item.component_property} for item in outputs],
            'inputs': [dict(item, value=values.get(item['id'], 1 if item['id'] == trigger else None)) for item in spec['inputs']],
            'state': [dict(item, value=values.get(item['id'] + '.' + item['property'], values.get(item['id']))) for item in spec['state']],
            'changedPropIds': [trigger + ('.contents' if trigger == 'qa-upload' else '.n_clicks')],
        }
        return self.client.post(PREFIX + '_dash-update-component', json=payload, headers=headers or {})

    def crud(self, trigger, extra=None):
        values = {'qa-filter': '', 'qa-record-id': 1, 'qa-record-version': 1,
                  'qa-delete-confirm': True, 'qa-upload': None}
        for field, value in zip(FIELDS, ['IC', 'SYN999', 'Synthetic Test', 'TW', 'Tainan', 'B', 'LEVEL 1']):
            values['qa-field-' + field] = value
        values.update(extra or {})
        # Browser workflow now queries the target before update/delete and
        # keeps create fields in a separate modal.
        values.update({'qa-create-field-' + f: values.get('qa-field-' + f) for f in FIELDS})
        if trigger in ('qa-update', 'qa-submit-update', 'qa-delete'):
            is_delete = trigger == 'qa-delete'
            target = values.get('qa-record-id', 1)
            loaded = self.call('qa-delete-record.children' if is_delete else 'qa-load-status.children',
                'qa-query-delete' if is_delete else 'qa-load-id',
                {'qa-delete-id' if is_delete else 'qa-record-id': target})
            loaded = loaded.get_json()['response']
            proof_id = 'qa-loaded-delete' if is_delete else 'qa-loaded-update'
            values[proof_id] = loaded.get(proof_id, {}).get('data')
            values['qa-delete-id'] = target
        response = self.call('qa-status.children', trigger, values)
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return response.get_json()['response']

    def test_layout_and_login_csrf(self):
        self.assertEqual(self.client.get(PREFIX + '_dash-layout').status_code, 200)
        self.assertEqual(self.client.get(PREFIX + '_dash-dependencies').status_code, 200)
        self.assertEqual(self.client.post(PREFIX + 'demo-login', data={'identity': 'developer'}).status_code, 403)
        self.assertEqual(self.client.get(PREFIX, headers={'Host': 'evil.invalid'}).status_code, 403)
        self.login('reader')
        data = self.crud('qa-read')
        self.assertEqual(len(data['qa-table']['data']), 2)
        self.assertIn('未完成', self.crud('qa-create')['qa-status']['children'])
        self.assertEqual(len(self.crud('qa-read')['qa-table']['data']), 2)

    def test_auth_and_cross_origin(self):
        self.assertEqual(self.crud('qa-read')['qa-table']['data'], [])
        self.login('outsider')
        self.assertNotIn('qa-download', self.crud('qa-export'))
        self.login('editor')
        response = self.call('qa-status.children', 'qa-read', headers={'Origin': 'https://evil.invalid'})
        self.assertEqual(response.status_code, 403)
        response = self.client.post(PREFIX + '_dash-update-component', data='not json')
        self.assertEqual(response.status_code, 415)

    def test_mutations_filter_export_and_stale_version(self):
        self.login('editor')
        data = self.crud('qa-create')
        self.assertEqual(len(data['qa-table']['data']), 3)
        row = next(row for row in data['qa-table']['data'] if row['Vendor_Code'] == 'SYN999')
        args = {'qa-record-id': row['id'], 'qa-record-version': row['version']}
        args['qa-field-Rev'] = 'C'
        staged = self.crud('qa-update', args)
        token = staged['qa-change-token']['data']
        self.assertIn('committed', self.crud('qa-confirm-update', dict(args, **{'qa-change-token': token}))['qa-status']['children'])
        self.assertIn('未完成', self.crud('qa-update', args)['qa-status']['children'])
        filtered = self.crud('qa-export', {'qa-filter': 'Synthetic Test'})
        content = base64.b64decode(filtered['qa-download']['data']['content']).decode('utf-8-sig')
        self.assertEqual(len(list(csv.DictReader(io.StringIO(content)))), 1)
        self.assertEqual(len(filtered['qa-table']['data']), 1)
        args['qa-record-version'] = 2
        self.assertEqual(len(self.crud('qa-delete', args)['qa-table']['data']), 2)

    def test_stage_preview_import_and_replay(self):
        self.login('editor')
        stream = io.StringIO()
        writer = csv.DictWriter(stream, fieldnames=['id', 'version'] + FIELDS)
        writer.writeheader()
        writer.writerow(dict(zip(FIELDS, ['IC', 'SYNCSV', 'Synthetic Import', 'TW', 'Tainan', 'A', 'LEVEL 2'])))
        encoded = 'data:text/csv;base64,' + base64.b64encode(stream.getvalue().encode()).decode()
        staged = self.crud('qa-upload', {'qa-upload': encoded, 'qa-upload.filename': 'fixture.csv'})
        token = staged['qa-stage-token']['data']
        self.assertEqual(json.loads(staged['qa-preview']['children'])['count'], 1)
        self.login('developer')
        self.assertIn('未完成', self.crud('qa-import', {'qa-stage-token': token})['qa-status']['children'])
        self.login('editor')
        result = self.crud('qa-import', {'qa-stage-token': token})
        self.assertEqual(len(result['qa-table']['data']), 3)
        self.assertIn('未完成', self.crud('qa-import', {'qa-stage-token': token})['qa-status']['children'])

    def test_mock_mail_settings_run_log(self):
        self.login('editor')
        values = {'qa-mail-to': 'qa@example.invalid', 'qa-run-key': 'one', 'qa-mail-version': 0,
                  'qa-mail-cadence': 'daily', 'qa-mail-time': '10:00', 'qa-mail-timezone': 'Asia/Taipei',
                  'qa-mail-enabled': True, 'qa-mail-delivery': 'accepted', 'qa-source-ok': True, 'qa-report-ok': True}
        response = self.call('qa-mail-result.children', 'qa-mail-save', values)
        self.assertEqual(response.status_code, 200)
        data = response.get_json()['response']
        self.assertEqual(data['qa-mail-version']['data'], 1)
        values['qa-mail-version'] = 1
        response = self.call('qa-mail-result.children', 'qa-mail-run', values)
        self.assertEqual(response.status_code, 200)
        self.assertIn('accepted', response.get_json()['response']['qa-mail-result']['children'])
        response = self.call('qa-mail-result.children', 'qa-mail-run', values)
        self.assertEqual(response.status_code, 200)
        self.assertIn('already claimed', response.get_json()['response']['qa-mail-result']['children'].lower())
        values['qa-run-key'] = 'source-failed'
        values['qa-source-ok'] = False
        response = self.call('qa-mail-result.children', 'qa-mail-run', values)
        self.assertIn('source_failed', response.get_json()['response']['qa-mail-result']['children'])
        response = self.call('qa-mail-result.children', 'qa-mail-refresh', values)
        self.assertIn('next_scheduled_times', response.get_json()['response']['qa-mail-result']['children'])
        self.login('reader')
        response = self.call('qa-mail-result.children', 'qa-mail-save', values)
        self.assertIn('未完成', response.get_json()['response']['qa-mail-result']['children'])

    def test_xlsx_export_reimport(self):
        from openpyxl import load_workbook
        self.login('editor')
        exported = self.crud('qa-export-xlsx', {'qa-filter': 'Vendor 1'})
        content = base64.b64decode(exported['qa-download']['data']['content'])
        workbook = load_workbook(io.BytesIO(content))
        self.assertEqual(workbook.active.max_row, 2)
        workbook.active.cell(2, 1).value = None
        workbook.active.cell(2, 2).value = None
        workbook.active.cell(2, 4).value = 'SYNXLSX'
        target = io.BytesIO()
        workbook.save(target)
        encoded = 'data:application/octet-stream;base64,' + base64.b64encode(target.getvalue()).decode()
        result = self.crud('qa-upload', {'qa-upload': encoded, 'qa-upload.filename': 'fixture.xlsx'})
        self.assertEqual(json.loads(result['qa-preview']['children'])['count'], 1)
        result = self.crud('qa-import', {'qa-stage-token': result['qa-stage-token']['data']})
        self.assertEqual(len(result['qa-table']['data']), 3)

    def test_admin_is_not_crud_and_oversize_request(self):
        self.login('admin')
        self.assertEqual(len(self.crud('qa-read')['qa-table']['data']), 2)
        self.assertIn('未完成', self.crud('qa-create')['qa-status']['children'])
        response = self.client.post(PREFIX + 'demo-login', data={'padding': 'x' * (16 * 1024 * 1024)})
        self.assertEqual(response.status_code, 413)

    def test_load_record_populates_form_with_server_version(self):
        self.login('editor')
        result = self.call('qa-load-status.children', 'qa-load-id', {'qa-record-id': 1})
        self.assertEqual(result.status_code, 200)
        data = result.get_json()['response']
        self.assertEqual(data['qa-field-Vendor_Code']['value'], 'SYN001')
        self.assertEqual(data['qa-record-version']['value'], 1)
        self.login('outsider')
        result = self.call('qa-load-status.children', 'qa-load-id', {'qa-record-id': 1})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.get_json()['response']['qa-field-Vendor_Code']['value'], '')
        self.assertIsNone(result.get_json()['response']['qa-loaded-update']['data'])


if __name__ == '__main__':
    unittest.main()
