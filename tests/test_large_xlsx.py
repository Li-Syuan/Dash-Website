"""Bounded XLSX snapshots preserve content, isolation and publication checks."""
from datetime import date, datetime
import io
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import openpyxl
from openpyxl.worksheet._write_only import WriteOnlyWorksheet
from openpyxl.worksheet._writer import ALL_TEMP_FILES

from reporting_workspace.legacy_crud import (
    CSV_FIELDS, QSL_FIELDS, InvalidInput, LegacyCrudService, PermissionDenied,
    StorageUnavailable,
)
from reporting_workspace.legacy_policy import Policy


def record(index, **changes):
    values = dict(Material_Type='Synthetic', Vendor_Code='SYN-{:06d}'.format(index),
                  Vendor_Name='Synthetic supplier', Country='TW', City='Synthetic city',
                  Rev='A', Supplier_Level='LEVEL 1')
    values.update(changes)
    return values


class LargeXlsxTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix='large-xlsx-contract-')
        self.path = str(Path(self.directory.name) / 'synthetic.sqlite')
        self.policy = Policy(orgcode=['SYNTHETIC'], crud_roles=['dev'])
        self.actor = SimpleNamespace(id='synthetic-owner', orgcode='SYNTHETIC-A',
                                    is_authenticated=True, is_dev=True, is_admin=False)
        self.services = []
        self.service = self.open_service(allow_upload=True)
        self.tempfiles_before = set(ALL_TEMP_FILES)

    def tearDown(self):
        for service in self.services:
            service.close()
        self.directory.cleanup()
        self.assertEqual(set(ALL_TEMP_FILES), self.tempfiles_before,
                         'Export retained an openpyxl temporary worksheet')

    def open_service(self, **options):
        service = LegacyCrudService(self.path, lambda actor: self.policy, **options)
        self.services.append(service)
        return service

    def seed(self, count):
        for start in range(0, count, 500):
            rows = [record(index) for index in range(start, min(count, start + 500))]
            outcome = self.service.submit_stage(self.actor, self.service.stage_rows(self.actor, rows))
            self.assertTrue(outcome.committed)
            self.assertEqual(outcome.failed, 0)

    def values(self, content):
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True,
                                          data_only=False, keep_links=False)
        try:
            self.assertEqual(workbook.sheetnames, ['QSL'])
            rows = list(workbook['QSL'].iter_rows())
            self.assertEqual(tuple(cell.value for cell in rows[0]), CSV_FIELDS)
            for row in rows[1:]:
                self.assertEqual([cell.data_type for cell in row], ['n', 'n'] + ['s'] * 7)
            return [tuple(cell.value for cell in row) for row in rows[1:]]
        finally:
            workbook.close()

    def test_multiple_batches_preserve_count_id_order_values_and_types(self):
        self.seed(2051)
        expected = self.service.query(self.actor, limit=3000)
        self.assertEqual(self.values(self.service.export_xlsx(self.actor)),
                         [tuple(row[key] for key in CSV_FIELDS) for row in expected])
        self.assertEqual(self.service.max_export_rows, 50000)

    def test_dates_numeric_looking_text_and_formula_error_literals_are_not_coerced(self):
        special = ['=1+1', '+001', '-001', '@SUM(A1)', '#N/A', '#DIV/0!', '00123',
                   '2026-10-04', '2026-10-04T12:34:56', '\u53f0\u5317\U0001f642', 'a\nb', 'A&B<x>']
        for index, value in enumerate(special):
            self.service.create(self.actor, record(index, Rev=value))
        rows = self.values(self.service.export_xlsx(self.actor))
        self.assertEqual([row[CSV_FIELDS.index('Rev')] for row in rows], special)
        for value in (date(2026, 10, 4), datetime(2026, 10, 4, 12, 34), 123, True):
            with self.subTest(value_type=type(value).__name__), self.assertRaises(InvalidInput):
                self.service.create(self.actor, record(999, Rev=value))

    def test_empty_optional_fields_keep_existing_blank_xlsx_roundtrip(self):
        self.service.create(self.actor, record(0, Rev='', Supplier_Level=''))
        content = self.service.export_xlsx(self.actor)
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            cells = list(workbook['QSL'].iter_rows())[1]
            self.assertIsNone(cells[CSV_FIELDS.index('Rev')].value)
            self.assertIsNone(cells[CSV_FIELDS.index('Supplier_Level')].value)
        finally:
            workbook.close()
        result = self.service.submit_stage(self.actor, self.service.stage_xlsx(self.actor, content))
        self.assertEqual(result.updated, 1)
        row = self.service.get(self.actor, 1)
        self.assertEqual((row['Rev'], row['Supplier_Level']), ('', ''))

    def test_filter_target_deleted_boundary_and_exact_cap_remain_enforced(self):
        for index in range(4):
            self.service.create(self.actor, record(index, Vendor_Name='Synthetic %_' if index < 2 else 'Other'))
        self.service.delete(self.actor, 4, 1)
        other = self.open_service(target='other-synthetic')
        other.create(self.actor, record(0, Vendor_Name='Synthetic %_'))
        bounded = self.open_service(max_export_rows=2)
        rows = self.values(bounded.export_xlsx(self.actor, {'Vendor_Name': '%_'}))
        self.assertEqual([row[0] for row in rows], [1, 2])
        with self.assertRaises(InvalidInput):
            bounded.export_xlsx(self.actor)
        for filters in ({'unknown': 'x'}, {'Vendor_Name': 1}, [], 'bad'):
            with self.subTest(filters=filters), self.assertRaises(InvalidInput):
                self.service.export_xlsx(self.actor, filters)

    def test_snapshot_stays_consistent_and_releases_database_before_workbook_writes(self):
        self.seed(1030)
        before = self.service.query(self.actor, limit=2000)
        writer = self.open_service()
        original = WriteOnlyWorksheet.append
        changed = []

        def append(sheet, row):
            if not changed:
                # A separate real connection must commit before serialization
                # continues; retaining the source lock would hit busy timeout.
                writer.update(self.actor, 1030, 1, {'Rev': 'B'})
                writer.delete(self.actor, 1, 1)
                writer.create(self.actor, record(9999))
                changed.append(True)
            return original(sheet, row)

        with patch.object(WriteOnlyWorksheet, 'append', append):
            content = self.service.export_xlsx(self.actor)
        self.assertEqual(self.values(content), [tuple(row[key] for key in CSV_FIELDS) for row in before])
        self.assertEqual(writer.get(self.actor, 1030)['Rev'], 'B')
        self.assertEqual(len(writer.query(self.actor, limit=2000)), 1030)

    def test_policy_revocation_during_serialization_prevents_publication_and_cleans_files(self):
        self.seed(1030)
        original = WriteOnlyWorksheet.append
        calls = []

        def append(sheet, row):
            result = original(sheet, row)
            calls.append(True)
            if len(calls) == 50:
                self.policy = Policy(orgcode=['DENIED'])
            return result

        with patch.object(WriteOnlyWorksheet, 'append', append), self.assertRaises(PermissionDenied):
            self.service.export_xlsx(self.actor)
        self.assertLess(len(calls), 1031)

    def test_actor_replacement_during_serialization_cannot_receive_original_snapshot(self):
        self.seed(4)
        original = WriteOnlyWorksheet.append

        def append(sheet, row):
            result = original(sheet, row)
            self.actor.id = 'different-synthetic-owner'
            return result

        with patch.object(WriteOnlyWorksheet, 'append', append), self.assertRaises(PermissionDenied):
            self.service.export_xlsx(self.actor)

    def test_policy_revocation_during_zip_save_is_checked_before_return(self):
        self.seed(4)
        original = openpyxl.Workbook.save

        def save(workbook, output):
            result = original(workbook, output)
            self.policy = Policy(orgcode=['DENIED'])
            return result

        with patch.object(openpyxl.Workbook, 'save', save), self.assertRaises(PermissionDenied):
            self.service.export_xlsx(self.actor)

    def test_revocation_before_first_row_does_not_create_an_orphan_worksheet(self):
        self.seed(4)
        original = openpyxl.Workbook.create_sheet

        def create_sheet(workbook, *args, **kwargs):
            sheet = original(workbook, *args, **kwargs)
            self.policy = Policy(orgcode=['DENIED'])
            return sheet

        with patch.object(openpyxl.Workbook, 'create_sheet', create_sheet):
            with self.assertRaises(PermissionDenied):
                self.service.export_xlsx(self.actor)
        self.assertEqual(set(ALL_TEMP_FILES), self.tempfiles_before)

    def test_archive_write_failure_is_sanitized_and_does_not_change_source(self):
        self.seed(4)
        before = self.service.query(self.actor)
        with patch.object(openpyxl.Workbook, 'save', side_effect=OSError('private path payload')):
            with self.assertRaises(StorageUnavailable) as raised:
                self.service.export_xlsx(self.actor)
        self.assertNotIn('private', str(raised.exception))
        self.assertEqual(self.service.query(self.actor), before)

    def test_snapshot_disk_failure_closes_its_private_file_without_writing_source(self):
        self.seed(4)
        before = self.service.query(self.actor)
        create_file = tempfile.TemporaryFile
        created = []

        class FailedSpool:
            def __init__(self, *args, **kwargs):
                self.file = create_file(*args, **kwargs)
                created.append(self.file)

            def __enter__(self):
                return self

            def __exit__(self, *args):
                self.file.close()

            def write(self, data):
                raise OSError('synthetic disk full private path')

        with patch('reporting_workspace.legacy_crud.tempfile.TemporaryFile', FailedSpool):
            with self.assertRaises(StorageUnavailable) as raised:
                self.service.export_xlsx(self.actor)
        self.assertEqual(len(created), 1)
        self.assertTrue(all(file.closed for file in created))
        self.assertNotIn('private', str(raised.exception))
        self.assertEqual(self.service.query(self.actor), before)

    def test_policy_revocation_during_snapshot_prevents_workbook_creation(self):
        self.seed(1030)
        create_file = tempfile.TemporaryFile
        policy_owner = self
        created = []

        class RevokingSpool:
            def __init__(self, *args, **kwargs):
                self.file = create_file(*args, **kwargs)
                self.writes = 0
                created.append(self.file)

            def __enter__(self):
                return self

            def __exit__(self, *args):
                self.file.close()

            def write(self, data):
                result = self.file.write(data)
                self.writes += 1
                if self.writes == 10:
                    policy_owner.policy = Policy(orgcode=['DENIED'])
                return result

        with patch('reporting_workspace.legacy_crud.tempfile.TemporaryFile', RevokingSpool):
            with patch('openpyxl.Workbook') as workbook, self.assertRaises(PermissionDenied):
                self.service.export_xlsx(self.actor)
            workbook.assert_not_called()
        self.assertTrue(all(file.closed for file in created))

    def test_anonymous_or_denied_actor_cannot_obtain_an_empty_or_nonempty_workbook(self):
        for count in (0, 1):
            if count:
                self.seed(count)
            for changes in ({'is_authenticated': False}, {'orgcode': 'DENIED', 'is_dev': False}):
                attributes = vars(self.actor).copy()
                attributes.update(changes)
                with self.subTest(rows=count, changes=changes), self.assertRaises(PermissionDenied):
                    self.service.export_xlsx(SimpleNamespace(**attributes))


if __name__ == '__main__':
    unittest.main()
